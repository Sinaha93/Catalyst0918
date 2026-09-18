from __future__ import annotations

import argparse
import hashlib
import json
import mimetypes
import re
import sqlite3
import sys
import urllib.parse
import zipfile
from datetime import datetime, timezone
from http import HTTPStatus
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
from pathlib import Path
from xml.etree import ElementTree as ET

import openpyxl


ROOT = Path(__file__).resolve().parent
INPUT_DIR = ROOT / "input"
DATA_DIR = ROOT / "data"
OUTPUT_DIR = ROOT / "outputs"
STATIC_DIR = ROOT / "static"
DB_PATH = DATA_DIR / "catalyst_closing.db"
ALLOWED_UPLOADS = {".xlsx", ".xls", ".csv", ".pptx"}
MAX_UPLOAD_BYTES = 100 * 1024 * 1024


def now_iso() -> str:
    return datetime.now(timezone.utc).astimezone().isoformat(timespec="seconds")


def ensure_directories() -> None:
    for directory in (INPUT_DIR, DATA_DIR, OUTPUT_DIR, STATIC_DIR):
        directory.mkdir(parents=True, exist_ok=True)


def connect() -> sqlite3.Connection:
    ensure_directories()
    conn = sqlite3.connect(DB_PATH)
    conn.row_factory = sqlite3.Row
    conn.execute("PRAGMA foreign_keys = ON")
    return conn


def init_db() -> None:
    with connect() as conn:
        conn.executescript(
            """
            CREATE TABLE IF NOT EXISTS source_files (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                relative_path TEXT NOT NULL UNIQUE,
                file_name TEXT NOT NULL,
                sha256 TEXT NOT NULL,
                size_bytes INTEGER NOT NULL,
                modified_at TEXT NOT NULL,
                source_type TEXT NOT NULL,
                detected_period TEXT,
                status TEXT NOT NULL,
                sheet_count INTEGER,
                formula_count INTEGER NOT NULL DEFAULT 0,
                external_formula_count INTEGER NOT NULL DEFAULT 0,
                external_links_json TEXT NOT NULL DEFAULT '[]',
                warnings_json TEXT NOT NULL DEFAULT '[]',
                metadata_json TEXT NOT NULL DEFAULT '{}',
                scanned_at TEXT NOT NULL
            );

            CREATE TABLE IF NOT EXISTS manual_values (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                value_type TEXT NOT NULL,
                part_number TEXT,
                customer TEXT,
                effective_date TEXT NOT NULL,
                numeric_value REAL,
                text_value TEXT,
                reason TEXT NOT NULL,
                created_at TEXT NOT NULL
            );

            CREATE TABLE IF NOT EXISTS import_profiles (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                profile_name TEXT NOT NULL UNIQUE,
                source_type TEXT NOT NULL,
                header_signature TEXT NOT NULL,
                mapping_json TEXT NOT NULL,
                created_at TEXT NOT NULL,
                updated_at TEXT NOT NULL
            );

            CREATE TABLE IF NOT EXISTS part_master (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                part_number TEXT NOT NULL UNIQUE,
                vehicle TEXT,
                supplier TEXT,
                customer TEXT,
                plant TEXT,
                active INTEGER NOT NULL DEFAULT 1,
                source_file_id INTEGER,
                updated_at TEXT NOT NULL,
                FOREIGN KEY(source_file_id) REFERENCES source_files(id)
            );

            CREATE TABLE IF NOT EXISTS price_history (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                part_number TEXT NOT NULL,
                price REAL NOT NULL,
                valid_from TEXT NOT NULL,
                valid_to TEXT,
                price_type TEXT,
                source_file_id INTEGER,
                created_at TEXT NOT NULL,
                FOREIGN KEY(source_file_id) REFERENCES source_files(id)
            );

            CREATE TABLE IF NOT EXISTS transactions (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                transaction_date TEXT NOT NULL,
                transaction_type TEXT NOT NULL,
                part_number TEXT NOT NULL,
                customer TEXT,
                plant TEXT,
                quantity REAL NOT NULL,
                amount REAL,
                source_file_id INTEGER NOT NULL,
                source_sheet TEXT,
                source_row INTEGER,
                created_at TEXT NOT NULL,
                FOREIGN KEY(source_file_id) REFERENCES source_files(id)
            );

            CREATE TABLE IF NOT EXISTS closing_runs (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                closing_period TEXT NOT NULL,
                status TEXT NOT NULL,
                source_snapshot_json TEXT NOT NULL,
                notes TEXT,
                created_at TEXT NOT NULL,
                completed_at TEXT
            );
            """
        )


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for block in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def relative_path(path: Path) -> str:
    try:
        return str(path.resolve().relative_to(ROOT)).replace("\\", "/")
    except ValueError:
        return str(path.resolve())


def detect_period(file_name: str, workbook=None) -> str | None:
    match = re.search(r"(?P<year>\d{2,4})\s*년\s*(?P<month>\d{1,2})\s*월", file_name)
    if match:
        year = int(match.group("year"))
        if year < 100:
            year += 2000
        month = int(match.group("month"))
        if 1 <= month <= 12:
            return f"{year:04d}-{month:02d}"
    if workbook is not None:
        for name in workbook.sheetnames:
            match = re.search(r"(?P<year>\d{2,4})\s*년\s*(?P<month>\d{1,2})\s*월", name)
            if match:
                year = int(match.group("year"))
                if year < 100:
                    year += 2000
                return f"{year:04d}-{int(match.group('month')):02d}"
    return None


def classify_workbook(sheet_names: list[str], file_name: str) -> str:
    names = set(sheet_names)
    if {"종합", "종합2"}.issubset(names):
        return "월마감 원본"
    if "납품 Summary" in names:
        return "발주·납품 계획"
    if "8월 마감자료" in names or {"Sheet2", "글로비스"}.issubset(names):
        return "출하실적 대사"
    lowered = file_name.lower()
    if "단가" in file_name:
        return "단가 마스터"
    if "촉매" in file_name and "현황" in file_name:
        return "품번 마스터"
    if "입고" in file_name:
        return "입고 실적"
    if "주문" in file_name or "발주" in file_name:
        return "발주 원본"
    if lowered.endswith(".csv"):
        return "CSV 입력"
    return "미분류 엑셀"


def workbook_external_targets(path: Path) -> list[str]:
    targets: list[str] = []
    try:
        with zipfile.ZipFile(path) as archive:
            rel_files = sorted(
                name
                for name in archive.namelist()
                if name.startswith("xl/externalLinks/_rels/") and name.endswith(".rels")
            )
            for rel_name in rel_files:
                root = ET.fromstring(archive.read(rel_name))
                for relation in root:
                    target = relation.attrib.get("Target")
                    if target and target not in targets:
                        targets.append(urllib.parse.unquote(target))
    except (zipfile.BadZipFile, ET.ParseError, OSError):
        pass
    return targets


def inspect_xlsx(path: Path) -> dict:
    warnings: list[str] = []
    workbook = openpyxl.load_workbook(path, data_only=False, read_only=False, keep_links=True)
    cached = openpyxl.load_workbook(path, data_only=True, read_only=False, keep_links=True)
    formula_count = 0
    external_formula_count = 0
    formula_errors: list[str] = []
    cached_errors: list[str] = []
    nonempty_rows = 0

    for sheet in workbook.worksheets:
        cached_sheet = cached[sheet.title]
        for row in sheet.iter_rows():
            row_has_value = False
            for cell in row:
                value = cell.value
                if value is not None:
                    row_has_value = True
                if isinstance(value, str) and value.startswith("="):
                    formula_count += 1
                    if "[" in value and "]" in value:
                        external_formula_count += 1
                    if any(token in value.upper() for token in ("#REF!", "#DIV/0!", "#VALUE!", "#NAME?")):
                        formula_errors.append(f"{sheet.title}!{cell.coordinate}")
                cached_value = cached_sheet[cell.coordinate].value
                if isinstance(cached_value, str) and cached_value.startswith("#"):
                    cached_errors.append(f"{sheet.title}!{cell.coordinate}")
            if row_has_value:
                nonempty_rows += 1

    source_type = classify_workbook(workbook.sheetnames, path.name)
    external_targets = workbook_external_targets(path)
    if external_targets:
        warnings.append(f"외부 파일 참조 {len(external_targets)}개")
    if formula_errors or cached_errors:
        warnings.append(f"수식 오류 {len(set(formula_errors + cached_errors))}개")
    if source_type == "미분류 엑셀":
        warnings.append("입력 유형 매핑 필요")

    required_sheets = {
        "월마감 원본": {"종합", "종합2", "종합3(누적)"},
        "발주·납품 계획": {"납품 Summary"},
        "출하실적 대사": {"8월 마감자료"},
    }.get(source_type, set())
    missing_sheets = sorted(required_sheets.difference(workbook.sheetnames))
    if missing_sheets:
        warnings.append("필수 시트 누락: " + ", ".join(missing_sheets))

    return {
        "source_type": source_type,
        "detected_period": detect_period(path.name, workbook),
        "status": "검토 필요" if warnings else "사용 가능",
        "sheet_count": len(workbook.sheetnames),
        "formula_count": formula_count,
        "external_formula_count": external_formula_count,
        "external_links": external_targets,
        "warnings": warnings,
        "metadata": {
            "sheets": workbook.sheetnames,
            "nonempty_rows": nonempty_rows,
            "defined_names": len(workbook.defined_names),
        },
    }


def inspect_pptx(path: Path) -> dict:
    warnings: list[str] = []
    slide_count = 0
    embedded_count = 0
    try:
        with zipfile.ZipFile(path) as archive:
            names = archive.namelist()
            slide_count = len(
                [
                    name
                    for name in names
                    if re.fullmatch(r"ppt/slides/slide\d+\.xml", name)
                ]
            )
            embedded_count = len([name for name in names if name.startswith("ppt/embeddings/")])
    except zipfile.BadZipFile:
        warnings.append("PowerPoint 파일을 읽을 수 없음")
    if embedded_count:
        warnings.append(f"내장 개체 {embedded_count}개는 생성 시 별도 처리 필요")
    return {
        "source_type": "마감 보고서",
        "detected_period": detect_period(path.name),
        "status": "검토 필요" if warnings else "사용 가능",
        "sheet_count": slide_count,
        "formula_count": 0,
        "external_formula_count": 0,
        "external_links": [],
        "warnings": warnings,
        "metadata": {"slides": slide_count, "embedded_objects": embedded_count},
    }


def inspect_csv(path: Path) -> dict:
    warnings = ["CSV 열 매핑 필요"]
    return {
        "source_type": "CSV 입력",
        "detected_period": detect_period(path.name),
        "status": "검토 필요",
        "sheet_count": 1,
        "formula_count": 0,
        "external_formula_count": 0,
        "external_links": [],
        "warnings": warnings,
        "metadata": {},
    }


def inspect_source(path: Path) -> dict:
    suffix = path.suffix.lower()
    try:
        if suffix == ".xlsx":
            return inspect_xlsx(path)
        if suffix == ".pptx":
            return inspect_pptx(path)
        if suffix == ".csv":
            return inspect_csv(path)
        return {
            "source_type": "구형 엑셀",
            "detected_period": detect_period(path.name),
            "status": "검토 필요",
            "sheet_count": None,
            "formula_count": 0,
            "external_formula_count": 0,
            "external_links": [],
            "warnings": [".xls 파일은 .xlsx 변환 후 자동 분석 가능"],
            "metadata": {},
        }
    except Exception as exc:  # Preserve the file and expose the exact import failure.
        return {
            "source_type": "읽기 실패",
            "detected_period": detect_period(path.name),
            "status": "검토 필요",
            "sheet_count": None,
            "formula_count": 0,
            "external_formula_count": 0,
            "external_links": [],
            "warnings": [f"파일 분석 실패: {type(exc).__name__}: {exc}"],
            "metadata": {},
        }


def discover_sources() -> list[Path]:
    sources = [path for path in ROOT.iterdir() if path.is_file() and path.suffix.lower() in ALLOWED_UPLOADS]
    if INPUT_DIR.exists():
        sources.extend(
            path
            for path in INPUT_DIR.rglob("*")
            if path.is_file() and path.suffix.lower() in ALLOWED_UPLOADS
        )
    return sorted(set(sources), key=lambda item: str(item).lower())


def scan_sources() -> list[dict]:
    init_db()
    results = []
    with connect() as conn:
        for path in discover_sources():
            stat = path.stat()
            details = inspect_source(path)
            record = {
                "relative_path": relative_path(path),
                "file_name": path.name,
                "sha256": sha256_file(path),
                "size_bytes": stat.st_size,
                "modified_at": datetime.fromtimestamp(stat.st_mtime).astimezone().isoformat(timespec="seconds"),
                **details,
                "scanned_at": now_iso(),
            }
            conn.execute(
                """
                INSERT INTO source_files (
                    relative_path, file_name, sha256, size_bytes, modified_at,
                    source_type, detected_period, status, sheet_count, formula_count,
                    external_formula_count, external_links_json, warnings_json,
                    metadata_json, scanned_at
                ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
                ON CONFLICT(relative_path) DO UPDATE SET
                    file_name=excluded.file_name,
                    sha256=excluded.sha256,
                    size_bytes=excluded.size_bytes,
                    modified_at=excluded.modified_at,
                    source_type=excluded.source_type,
                    detected_period=excluded.detected_period,
                    status=excluded.status,
                    sheet_count=excluded.sheet_count,
                    formula_count=excluded.formula_count,
                    external_formula_count=excluded.external_formula_count,
                    external_links_json=excluded.external_links_json,
                    warnings_json=excluded.warnings_json,
                    metadata_json=excluded.metadata_json,
                    scanned_at=excluded.scanned_at
                """,
                (
                    record["relative_path"], record["file_name"], record["sha256"],
                    record["size_bytes"], record["modified_at"], record["source_type"],
                    record["detected_period"], record["status"], record["sheet_count"],
                    record["formula_count"], record["external_formula_count"],
                    json.dumps(record["external_links"], ensure_ascii=False),
                    json.dumps(record["warnings"], ensure_ascii=False),
                    json.dumps(record["metadata"], ensure_ascii=False), record["scanned_at"],
                ),
            )
            results.append(record)
    return results


def rows_to_dicts(rows) -> list[dict]:
    return [dict(row) for row in rows]


def list_sources() -> list[dict]:
    init_db()
    with connect() as conn:
        rows = conn.execute(
            "SELECT * FROM source_files ORDER BY detected_period DESC, file_name"
        ).fetchall()
    result = []
    for row in rows_to_dicts(rows):
        row["external_links"] = json.loads(row.pop("external_links_json"))
        row["warnings"] = json.loads(row.pop("warnings_json"))
        row["metadata"] = json.loads(row.pop("metadata_json"))
        result.append(row)
    return result


def dashboard_status() -> dict:
    init_db()
    with connect() as conn:
        source_count = conn.execute("SELECT COUNT(*) FROM source_files").fetchone()[0]
        ready_count = conn.execute("SELECT COUNT(*) FROM source_files WHERE status='사용 가능'").fetchone()[0]
        review_count = conn.execute("SELECT COUNT(*) FROM source_files WHERE status='검토 필요'").fetchone()[0]
        manual_count = conn.execute("SELECT COUNT(*) FROM manual_values").fetchone()[0]
        last_scan = conn.execute("SELECT MAX(scanned_at) FROM source_files").fetchone()[0]
    return {
        "source_count": source_count,
        "ready_count": ready_count,
        "review_count": review_count,
        "manual_count": manual_count,
        "last_scan": last_scan,
        "database": relative_path(DB_PATH),
        "automation_stage": "입력 파일 수집·검증",
    }


def add_manual_value(payload: dict) -> dict:
    value_type = str(payload.get("value_type", "")).strip()
    effective_date = str(payload.get("effective_date", "")).strip()
    reason = str(payload.get("reason", "")).strip()
    if not value_type or not effective_date or not reason:
        raise ValueError("구분, 적용일, 사유는 필수입니다.")
    numeric_value = payload.get("numeric_value")
    if numeric_value not in (None, ""):
        numeric_value = float(numeric_value)
    else:
        numeric_value = None
    created_at = now_iso()
    with connect() as conn:
        cursor = conn.execute(
            """
            INSERT INTO manual_values (
                value_type, part_number, customer, effective_date,
                numeric_value, text_value, reason, created_at
            ) VALUES (?, ?, ?, ?, ?, ?, ?, ?)
            """,
            (
                value_type,
                str(payload.get("part_number", "")).strip() or None,
                str(payload.get("customer", "")).strip() or None,
                effective_date,
                numeric_value,
                str(payload.get("text_value", "")).strip() or None,
                reason,
                created_at,
            ),
        )
        row_id = cursor.lastrowid
        row = conn.execute("SELECT * FROM manual_values WHERE id=?", (row_id,)).fetchone()
    return dict(row)


def list_manual_values() -> list[dict]:
    init_db()
    with connect() as conn:
        return rows_to_dicts(
            conn.execute("SELECT * FROM manual_values ORDER BY id DESC LIMIT 100").fetchall()
        )


def safe_upload_name(raw_name: str) -> str:
    decoded = urllib.parse.unquote(raw_name)
    name = Path(decoded).name.strip()
    if not name or Path(name).suffix.lower() not in ALLOWED_UPLOADS:
        raise ValueError("지원하는 파일 형식은 xlsx, xls, csv, pptx입니다.")
    return name


class AppHandler(BaseHTTPRequestHandler):
    server_version = "CatalystClosing/0.1"

    def log_message(self, format, *args):
        sys.stdout.write(f"[{self.log_date_time_string()}] {format % args}\n")

    def send_json(self, payload, status=HTTPStatus.OK):
        body = json.dumps(payload, ensure_ascii=False).encode("utf-8")
        self.send_response(status)
        self.send_header("Content-Type", "application/json; charset=utf-8")
        self.send_header("Content-Length", str(len(body)))
        self.end_headers()
        self.wfile.write(body)

    def read_json(self):
        length = int(self.headers.get("Content-Length", "0"))
        if length > 1024 * 1024:
            raise ValueError("요청 데이터가 너무 큽니다.")
        data = self.rfile.read(length) if length else b"{}"
        return json.loads(data.decode("utf-8"))

    def do_GET(self):
        parsed = urllib.parse.urlparse(self.path)
        if parsed.path == "/api/status":
            return self.send_json(dashboard_status())
        if parsed.path == "/api/sources":
            return self.send_json({"items": list_sources()})
        if parsed.path == "/api/manual-values":
            return self.send_json({"items": list_manual_values()})
        if parsed.path.startswith("/api/"):
            return self.send_json({"error": "API를 찾을 수 없습니다."}, HTTPStatus.NOT_FOUND)
        return self.serve_static(parsed.path)

    def do_POST(self):
        parsed = urllib.parse.urlparse(self.path)
        try:
            if parsed.path == "/api/scan":
                items = scan_sources()
                return self.send_json({"items": items, "status": dashboard_status()})
            if parsed.path == "/api/manual-values":
                item = add_manual_value(self.read_json())
                return self.send_json({"item": item}, HTTPStatus.CREATED)
            if parsed.path == "/api/upload":
                query = urllib.parse.parse_qs(parsed.query)
                file_name = safe_upload_name(query.get("filename", [""])[0])
                length = int(self.headers.get("Content-Length", "0"))
                if length <= 0 or length > MAX_UPLOAD_BYTES:
                    raise ValueError("파일 크기는 1바이트 이상 100MB 이하여야 합니다.")
                content = self.rfile.read(length)
                destination = INPUT_DIR / file_name
                if destination.exists():
                    stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
                    destination = INPUT_DIR / f"{destination.stem}_{stamp}{destination.suffix}"
                destination.write_bytes(content)
                return self.send_json(
                    {"file": relative_path(destination), "scan": scan_sources()},
                    HTTPStatus.CREATED,
                )
            return self.send_json({"error": "API를 찾을 수 없습니다."}, HTTPStatus.NOT_FOUND)
        except (ValueError, json.JSONDecodeError) as exc:
            return self.send_json({"error": str(exc)}, HTTPStatus.BAD_REQUEST)
        except Exception as exc:
            return self.send_json(
                {"error": f"서버 오류: {type(exc).__name__}: {exc}"},
                HTTPStatus.INTERNAL_SERVER_ERROR,
            )

    def serve_static(self, url_path: str):
        relative = "index.html" if url_path in ("", "/") else url_path.lstrip("/")
        requested = (STATIC_DIR / relative).resolve()
        try:
            requested.relative_to(STATIC_DIR.resolve())
        except ValueError:
            return self.send_error(HTTPStatus.FORBIDDEN)
        if not requested.is_file():
            requested = STATIC_DIR / "index.html"
        body = requested.read_bytes()
        content_type = mimetypes.guess_type(requested.name)[0] or "application/octet-stream"
        if content_type.startswith("text/"):
            content_type += "; charset=utf-8"
        self.send_response(HTTPStatus.OK)
        self.send_header("Content-Type", content_type)
        self.send_header("Content-Length", str(len(body)))
        self.end_headers()
        self.wfile.write(body)


def main() -> None:
    parser = argparse.ArgumentParser(description="촉매 마감 자동화 로컬 앱")
    parser.add_argument("--host", default="127.0.0.1")
    parser.add_argument("--port", type=int, default=8765)
    parser.add_argument("--scan-once", action="store_true", help="파일을 스캔하고 JSON 결과를 출력")
    args = parser.parse_args()
    init_db()
    if args.scan_once:
        print(json.dumps({"items": scan_sources(), "status": dashboard_status()}, ensure_ascii=False, indent=2))
        return
    server = ThreadingHTTPServer((args.host, args.port), AppHandler)
    print(f"촉매 마감 자동화: http://{args.host}:{args.port}")
    try:
        server.serve_forever()
    except KeyboardInterrupt:
        print("\n서버를 종료합니다.")
    finally:
        server.server_close()


if __name__ == "__main__":
    main()
