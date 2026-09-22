"""Office template exporter. Workbooks are rendered by the installed Excel engine."""
import json
import shutil
import tempfile
import hashlib
from pathlib import Path
from datetime import date
from decimal import Decimal
import openpyxl
from . import store

CUSTOMERS=['기아 화성','기아 광명','기아 광주','현대 아산','현대 울산','현대 전주','현대 글로비스','현대 위아','세종공업','현대 모비스']

def template(suffix):
    matches=[p for p in store.ROOT.glob('*'+suffix) if '촉매 마감' in p.name]
    if not matches:
        bundled=Path(__file__).resolve().parents[1]/'templates'
        matches=list(bundled.glob('*'+suffix))
    if not matches: raise ValueError('기존 촉매 마감 '+suffix+' 양식이 필요합니다')
    return sorted(matches)[0]

def freeze_context(period):
    context={'templates':{},'earlier':[]}
    for suffix in ('.xlsx',):
        try:source=template(suffix)
        except ValueError:continue
        digest=hashlib.sha256(source.read_bytes()).hexdigest()
        target=store.ROOT/'archive'/(digest+suffix)
        if not target.exists():shutil.copy2(source,target)
        context['templates'][suffix]={'hash':digest,'file':target.name,'name':source.name}
    with store.db() as c:
        for row in c.execute('SELECT snapshot FROM runs WHERE period<? ORDER BY period,version',(period,)):
            s=json.loads(row[0])
            context['earlier'].append({'period':s['period'],'customers':s['customers']})
    return context

def capture_range(excel,ws,address,path):
    ws.Activate()
    rng=ws.Range(address)
    rng.CopyPicture(Appearance=1,Format=2)
    chart=ws.ChartObjects().Add(0,0,rng.Width,rng.Height)
    try:
        chart.Activate()
        chart.Chart.Paste()
        if not chart.Chart.Export(str(path),'PNG'):
            raise RuntimeError('Excel 범위 이미지 저장 실패')
    finally:
        chart.Delete()
        excel.CutCopyMode=False

def write_cell(ws,r,c,value):
    ws.Cells(r,c).Value=float(value) if isinstance(value,(Decimal,int,float)) else value

def simplify_trend_sheet(ws):
    """Leave the trend worksheet as a chart-only report view."""
    ws.Range('B66:AT148').ClearContents()
    ws.Rows('30:148').Hidden=True
    charts=ws.ChartObjects()
    for index in range(1,charts.Count+1):
        charts.Item(index).Chart.PlotVisibleOnly=False

def generate(rid):
    import pythoncom
    import win32com.client
    with store.db() as c:
        run=c.execute('SELECT * FROM runs WHERE id=?',(rid,)).fetchone()
    if not run: raise ValueError('확정 마감을 찾을 수 없습니다')
    snap=json.loads(run['snapshot'])
    from .report_estimates import report_note
    for row in snap['details']:
        row['note']=report_note(row)
    unknown=set(r['customer'] for r in snap['details'])-set(CUSTOMERS)
    if unknown: raise ValueError('기존 양식에 없는 거래처입니다: '+', '.join(unknown))
    period=run['period']; year,month=map(int,period.split('-'))
    stem=f'{year}년 {month}월 촉매 마감_v{run["version"]}'
    context=snap.get('report_context')
    if not context or '.xlsx' not in context.get('templates',{}):
        raise ValueError('확정 시점의 보고서 양식이 없습니다. 양식을 등록한 후 수정 마감을 확정해주세요.')
    def frozen(suffix):
        info=context['templates'][suffix]; path=store.ROOT/'archive'/info['file']
        if not path.exists() or hashlib.sha256(path.read_bytes()).hexdigest()!=info['hash']:
            raise ValueError('보관된 보고서 양식이 없거나 변경되었습니다')
        return path
    xsource=frozen('.xlsx')
    baseline=openpyxl.load_workbook(xsource,data_only=True)
    # Never publish a stale month series under a new reporting date.
    work=Path(tempfile.mkdtemp(prefix='catalyst_report_'))
    xout=work/(stem+'.xlsx')
    shutil.copy2(xsource,xout)
    excel=None; workbook=None
    pythoncom.CoInitialize()
    try:
        excel=win32com.client.DispatchEx('Excel.Application')
        excel.Visible=False
        excel.DisplayAlerts=False
        excel.AskToUpdateLinks=False
        excel.AutomationSecurity=3
        workbook=excel.Workbooks.Open(str(xout),UpdateLinks=0,ReadOnly=False)
        links=workbook.LinkSources(1)
        if links:
            for link in links: workbook.BreakLink(link,1)
        customer_totals={r['customer']:r for r in snap['customers']}
        detail_ranges={}
        history=[]
        cumulative=baseline['종합3(누적)']
        for cells in cumulative.iter_rows(min_row=2,values_only=True):
            if cells[0] and cells[1] and cells[2] and cells[5] is not None:
                history.append({'year':int(cells[0]),'month':int(cells[1]),'customer':cells[2],'type':cells[3],'metric':cells[4],'value':cells[5]})
        history_map={(r['year'],r['month'],r['customer'],r['type'],r['metric']):r['value'] for r in history}
        for s in context['earlier']:
            y,m=map(int,s['period'].split('-'))
            for r in s['customers']:
                for typ,metric,field in [('실적','수량','settlement'),('실적','금액','settlement_amount'),('계획','수량','plan')]:
                    history_map[(y,m,r['customer'],typ,metric)]=float(r[field]) if r[field] is not None else None
        for name in CUSTOMERS:
            ws=workbook.Worksheets(name); source=baseline[name]
            data=[r for r in snap['details'] if r['customer']==name and Decimal(r['settlement'])!=0]
            totalrow=next((r for r in range(9,source.max_row+1) if str(source.cell(r,2).value).strip()=='TOTAL'),None)
            if totalrow is None: raise ValueError(name+' 계산서 합계행을 찾을 수 없습니다')
            capacity=totalrow-9
            if len(data)>capacity:
                extra=len(data)-capacity
                ws.Range(f'B{totalrow}:H{totalrow+extra-1}').Insert(Shift=-4121)
                totalrow+=extra
            ws.Range(f'B9:H{totalrow-1}').ClearContents()
            for i,r in enumerate(data,9):
                q=Decimal(r['settlement']); amount=Decimal(r['settlement_amount'])
                values=[i-8,r['part'],float(amount/q),float(q),float(amount),r.get('supplier',''),r['note']]
                ws.Range(f'B{i}:H{i}').Value=(tuple(values),)
            ws.Cells(totalrow,2).Value='TOTAL'
            ws.Cells(totalrow,5).Formula=f'=SUM(E9:E{totalrow-1})'
            ws.Cells(totalrow,6).Formula=f'=SUM(F9:F{totalrow-1})'
            ws.Range('B2').Value=f'{month}월 촉매 계산서'
            ws.Range('B5').Value=f'작성일 : {date.today().isoformat()}'
            detail_ranges[name]=f'B2:H{totalrow}'
            current=customer_totals.get(name,{})
            for typ,metric,key in [('실적','수량','settlement'),('실적','금액','settlement_amount'),('계획','수량','plan')]:
                value=current.get(key,None if key=='plan' else 0)
                history_map[(year,month,name,typ,metric)]=float(value) if value is not None else None
            plan_records=[r for r in snap['input_snapshot'] if r['kind']=='plan' and r['customer']==name and r['date'].startswith(period)]
            plan_amount=sum(Decimal(r['amount']) for r in plan_records if r['amount'] is not None) if plan_records and all(r['amount'] is not None for r in plan_records) else None
            history_map[(year,month,name,'계획','금액')]=float(plan_amount) if plan_amount is not None else None
            header=next(r for r in range(1,source.max_row+1) if source.cell(r,10).value=='구분')
            update_history(ws,header,10,name,history_map,year,month)
        summary=workbook.Worksheets('종합')
        summary.Range('B2').Value=f'※ {month}월 촉매 매입 및 미정산 품목 (종합)'
        for col,text in [(6,f'{month}月 입고\n(B)'),(8,f'{month}月 계산서\n(C)'),(10,f'{month}月 미결\n(A+B-C)')]:summary.Cells(4,col).Value=text
        summary.Range('B19').Value=f'※ {month}월 미결사유 및 특이사항'
        for col,window in ((12,3),(13,6),(14,12)):
            a=year*12+month-1-window; b=year*12+month-2
            summary.Cells(6,col).Value=f'{window}개월\n({a//12%100}.{a%12+1}~{b//12%100}.{b%12+1})'
        summary.Cells(6,15).Value=f'{(year-1)%100}년'
        for idx,name in enumerate(CUSTOMERS,7):
            rows=[r for r in snap['details'] if r['customer']==name]
            sums={k:sum(Decimal(r[k]) for r in rows) for k in ('opening','opening_amount','receipt','receipt_amount','settlement','settlement_amount','closing','closing_amount')}
            summary.Range(f'D{idx}:K{idx}').Value=(tuple(float(sums[k]) for k in ('opening','opening_amount','receipt','receipt_amount','settlement','settlement_amount','closing','closing_amount')),)
            summary.Cells(idx,15).Value=historical_average(history_map,name,year-1,'실적','수량')
        # Replace the unresolved-item content within the existing formatted worksheet.
        pending=workbook.Worksheets('미정산품목(26년6월)')
        pending.Range('B1').Value=f'▣ {year%100}년 {month}월 이월 및 정산품목'
        pending.Range('I3').Value=f'{month}월 매입'; pending.Range('J3').Value=f'{month}월 마감'
        from .pending_report import grouped_rows, write_excel as write_pending
        write_pending(pending,grouped_rows(snap['details'],CUSTOMERS))
        from .report_estimates import write_excel
        write_excel(pending,snap)
        # Existing chart series retain their source ranges; the 13-month data window is rolled.
        trend_name='월별 계획·실적' if '월별 계획·실적' in baseline.sheetnames else '종합2'
        combined=workbook.Worksheets(trend_name)
        if trend_name=='종합2':combined.Name='월별 계획·실적'
        update_history(combined,30,2,None,history_map,year,month)
        simplify_trend_sheet(combined)
        cum=workbook.Worksheets('종합3(누적)')
        cum.UsedRange.ClearContents()
        cum.Range('A1:F1').Value=(('연도','월','거래처','구분','항목','값'),)
        history_rows=[(*key,value) for key,value in sorted(history_map.items()) if value is not None]
        if history_rows:cum.Range(f'A2:F{len(history_rows)+1}').Value=tuple(history_rows)
        excel.CalculateFullRebuild()
        workbook.Save()
        workbook.Close(SaveChanges=True);workbook=None
        verify_output(xout,snap)
        shutil.copy2(xout,store.ROOT/'outputs'/xout.name)
        return {'files':[xout.name]}
    finally:
        if workbook is not None:workbook.Close(SaveChanges=False)
        if excel is not None:excel.Quit()
        pythoncom.CoUninitialize()

def verify_output(xout,snapshot):
    """Validate the calculated workbook before publication; no PowerPoint needed."""
    wb=openpyxl.load_workbook(xout,data_only=True,read_only=True)
    try:
        if len(wb.sheetnames)!=14:raise ValueError('출력 엑셀 시트 수가 기준과 다릅니다')
        expected=sum(Decimal(r['settlement_amount']) for r in snapshot['customers'])
        actual=Decimal(str(wb['종합']['I17'].value or 0))
        if abs(expected-actual)>Decimal('0.01'):raise ValueError('출력 엑셀 종합 금액이 확정 결과와 다릅니다')
        from .report_estimates import estimate_rows
        from .pending_report import grouped_rows
        for index,row in enumerate(grouped_rows(snapshot['details'],CUSTOMERS),4):
            if not row['subtotal']:continue
            for col,key in ((8,'opening'),(9,'receipt'),(10,'settlement'),(11,'closing'),(12,'closing_amount')):
                value=wb['미정산품목(26년6월)'].cell(index,col).value
                if value is None or abs(Decimal(str(value))-Decimal(row[key]))>Decimal('0.01'):
                    raise ValueError('미정산 납품처 소계 대사 불일치: '+row['customer'])
        for index,row in enumerate(estimate_rows(snapshot),5):
            value=wb['미정산품목(26년6월)'].cell(index,21).value
            if value is None or abs(Decimal(str(value))-row['estimated_amount'])>Decimal('0.01'):
                raise ValueError('엑셀 예상 정산금액 대사 불일치')
        for sheet in wb:
            for row in sheet:
                if any(c.data_type=='e' for c in row):
                    raise ValueError(sheet.title+' 시트에 Excel 수식 오류가 있습니다')
    finally:wb.close()

def update_history(ws,header,col,customer,history,year,month):
    metrics=[('계획','수량'),('실적','수량'),('계획','금액'),('실적','금액')]
    for offset in range(13):
        absolute=year*12+month-1-12+offset
        y,m=absolute//12,absolute%12+1
        ws.Cells(header,col+1+offset).Value=f'{y%100}년 {m}월'
        for idx,(typ,metric) in enumerate(metrics,1):
            if customer:
                value=history.get((y,m,customer,typ,metric))
            else:
                values=[history.get((y,m,c,typ,metric)) for c in CUSTOMERS]
                value=sum(values) if all(v is not None for v in values) else None
            ws.Cells(header+idx,col+1+offset).Value=value
    for idx,(typ,metric) in enumerate(metrics,1):
        for j,window in enumerate((3,6,12),14):
            values=[]
            for offset in range(window):
                absolute=year*12+month-2-offset
                keys=[customer] if customer else CUSTOMERS
                months=[history.get((absolute//12,absolute%12+1,c,typ,metric)) for c in keys]
                values.append(sum(months) if all(v is not None for v in months) else None)
            ws.Cells(header+idx,col+j).Value=sum(values)/window if all(v is not None for v in values) else None
    for idx,metric in ((5,'수량'),(6,'금액')):
        ws.Cells(header+idx,col).Value=f'{(year-1)%100}년 평균 {metric}'
        value=historical_average(history,customer,year-1,'실적',metric)
        for offset in range(1,17):ws.Cells(header+idx,col+offset).Value=value

def historical_average(history,customer,year,typ,metric):
    keys=[customer] if customer else CUSTOMERS
    values=[history.get((year,m,c,typ,metric)) for m in range(1,13) for c in keys]
    return sum(values)/12 if all(v is not None for v in values) else None
