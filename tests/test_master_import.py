import json
from pathlib import Path
import openpyxl
import pytest
from catalyst import store,imports,master_import,closing,source_lifecycle


@pytest.fixture
def setup(tmp_path,monkeypatch):
    monkeypatch.setattr(store,'ROOT',tmp_path);store.init()
    wb=openpyxl.Workbook();s=wb.active;s.title='BOM_마스터'
    for _ in range(8):s.append(['header'])
    s.cell(8,3,'촉매 마감처');s.cell(8,6,'촉매 사용구분');s.cell(8,13,'촉매 입력품번')
    s.append(['Customer','Parent','현대 전주','28991-4C520',1,'사용',None,None,None,None,None,'MParent','R60028991-4C520'])
    s=wb.create_sheet('촉매사_마스터')
    for _ in range(8):s.append(['header'])
    s.cell(8,1,'촉매 입력품번');s.cell(8,2,'촉매사')
    s.append(['R60028991-4C520','한국유미코아촉매'])
    path=tmp_path/'bom.xlsx';wb.save(path);wb.close();bom=imports.register(path)['id']
    wb=openpyxl.Workbook();s=wb.active
    s.append(['구매거래처','품번','통화','단가','적용종료일','입력일'])
    s.append(['Not the supplier','R60028991-4C520','KRW',177000,'9999-12-31','2026-09-21'])
    s.append(['Another buyer','28991-4C520','KRW',177000,'9999-12-31','2026-09-21'])
    path=tmp_path/'purchase_20260921.xlsx';wb.save(path);wb.close();price=imports.register(path)['id']
    return {'bom_id':bom,'price_id':price,'effective_date':'2026-08-01','reason':'확인'}


def test_master_pair_preview_apply_cancel_restore(setup):
    r=master_import.preview(setup)
    assert not r['errors'] and r['prices']==1 and r['links']==1
    assert r['rows'][-1]['provenance']['link']['supplier']=='한국유미코아촉매'
    assert r['rows'][1]['date']=='2026-08-01'
    master_import.apply({**setup,'fingerprint':r['fingerprint']})
    with pytest.raises(ValueError,match='이미 반영'):master_import.apply({**setup,'fingerprint':master_import.preview(setup)['fingerprint']})
    imports.commit_rows([{'kind':'receipt','part':'28991-4C520','customer':'현대 전주','date':'2026-08-31','quantity':2}],None,'receipt','append')
    row=closing.preview('2026-08')['details'][0]
    assert row['receipt_amount']=='354000' and row['supplier']=='한국유미코아촉매'
    sid=setup['price_id'];view=source_lifecycle.impact(sid)
    assert len(view['sources'])==2
    source_lifecycle.change(sid,{'fingerprint':view['fingerprint'],'reason':'취소'})
    assert closing.preview('2026-08')['details'][0]['receipt_amount'] is None
    source_lifecycle.change(sid,{'fingerprint':source_lifecycle.impact(sid)['fingerprint'],'reason':'복원'})
    assert closing.preview('2026-08')['details'][0]['receipt_amount']=='354000'


def test_effective_date_required_and_normalization(setup):
    assert master_import.part('R60028957-2JDF0_CKD')=='28957-2JDF0'
    assert master_import.part('28957-2JDF0_US')=='28957-2JDF0_US'
    with pytest.raises(ValueError):master_import.preview({**setup,'effective_date':''})


def test_missing_and_conflicting_prices_block_apply(setup):
    with store.db() as c:s=c.execute('SELECT * FROM sources WHERE id=?',(setup['price_id'],)).fetchone()
    path=imports.source_path(s)
    wb=openpyxl.load_workbook(path);wb.active.cell(3,4,1);wb.save(path);wb.close()
    r=master_import.preview(setup)
    assert r['errors']
    with pytest.raises(ValueError):master_import.apply({**setup,'fingerprint':r['fingerprint']})
